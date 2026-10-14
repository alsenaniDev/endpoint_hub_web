from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

OUT = Path(__file__).resolve().parents[1] / "docs" / "TWK_Endpoints_Hub_Executive_Overview.pptx"

NAVY = RGBColor(13, 28, 48)
INK = RGBColor(24, 39, 58)
MUTED = RGBColor(91, 108, 126)
PAPER = RGBColor(247, 249, 251)
WHITE = RGBColor(255, 255, 255)
CYAN = RGBColor(37, 184, 196)
CORAL = RGBColor(242, 111, 92)
MINT = RGBColor(117, 205, 165)
GOLD = RGBColor(242, 190, 76)
PALE_CYAN = RGBColor(225, 247, 248)
PALE_CORAL = RGBColor(255, 235, 231)
PALE_MINT = RGBColor(229, 247, 237)
PALE_GOLD = RGBColor(255, 246, 218)
FONT = "Aptos"
DISPLAY = "Aptos Display"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rgb(fill, transparency=None):
    return fill


def shape(slide, kind, x, y, w, h, fill=WHITE, line=None, radius=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line or fill
    if transparency is not None:
        s.fill.transparency = transparency
    return s


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.5, arrow=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    if arrow:
        ln.line.end_arrowhead = True
    return ln


def card(slide, x, y, w, h, title, body="", accent=CYAN, fill=WHITE, title_size=16, body_size=11):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, fill)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.07, h, accent, accent)
    text(slide, title, x + 0.22, y + 0.18, w - 0.38, 0.34, title_size, INK, True, DISPLAY)
    if body:
        text(slide, body, x + 0.22, y + 0.62, w - 0.38, h - 0.73, body_size, MUTED)
    return slide


def pill(slide, label, x, y, w, fill=PALE_CYAN, color=INK):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.32, fill, fill)
    text(slide, label, x, y + 0.055, w, 0.2, 9, color, True, FONT, PP_ALIGN.CENTER)


def add_notes(slide, message):
    notes = slide.notes_slide
    tf = notes.notes_text_frame
    tf.text = message


def base(title, kicker, number, dark=False):
    slide = prs.slides.add_slide(blank)
    bg = NAVY if dark else PAPER
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, bg, bg)
    if not dark:
        shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.1, CYAN, CYAN)
    text(slide, kicker.upper(), 0.55, 0.35, 5.5, 0.22, 9, CYAN if dark else CORAL, True, FONT)
    text(slide, title, 0.55, 0.67, 12.1, 0.54, 26, WHITE if dark else NAVY, True, DISPLAY)
    text(slide, f"TWK ENDPOINTS HUB  /  {number:02d}", 11.25, 7.05, 1.55, 0.2, 8, RGBColor(163, 178, 193) if dark else MUTED, True, FONT, PP_ALIGN.RIGHT)
    return slide


def big_metric(slide, value, label, x, y, accent=CYAN):
    text(slide, value, x, y, 1.5, 0.54, 28, accent, True, DISPLAY)
    text(slide, label, x, y + 0.5, 1.7, 0.4, 10, MUTED)


def flow_node(slide, label, x, y, w, h, fill=WHITE, accent=CYAN, size=12):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, fill)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, 0.06, accent, accent)
    text(slide, label, x + 0.08, y + 0.16, w - 0.16, h - 0.2, size, INK, True, FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# 1
s = base("A clearer way to work with APIs", "Executive overview", 1, True)
text(s, "TWK Endpoints Hub", 0.65, 1.8, 7.2, 0.72, 40, WHITE, True, DISPLAY)
text(s, "Discover. Test. Track. Govern.", 0.68, 2.65, 6.4, 0.42, 22, CYAN, True, DISPLAY)
text(s, "A centralized platform for turning scattered API knowledge into a shared, usable operating view.", 0.7, 3.25, 5.8, 0.8, 15, RGBColor(207, 220, 229))
for i, (label, color) in enumerate([("DISCOVER", CYAN), ("TEST", MINT), ("TRACK", GOLD), ("GOVERN", CORAL)]):
    pill(s, label, 0.7 + i * 1.52, 4.55, 1.32, color, NAVY)
shape(s, MSO_SHAPE.ARC, 8.6, 1.55, 3.4, 3.4, NAVY, CYAN)
for i, (label, y, color) in enumerate([("Teams", 2.0, CYAN), ("Systems", 2.85, MINT), ("Endpoints", 3.7, GOLD), ("Standards", 4.55, CORAL)]):
    flow_node(s, label, 9.35, y, 2.0, 0.55, NAVY, color, 13)
add_notes(s, "Say: TWK Endpoints Hub brings API discovery, testing, change visibility, sharing, and governance into one working environment.\nKey message: the product converts fragmented API knowledge into an operational platform.\nProblem addressed: teams lose time navigating scattered documentation and tools.\nBenefit to emphasize: one place to find and act on API information.")

# 2
s = base("API work is fragmented before it becomes manageable", "The problem", 2)
text(s, "The cost is not only technical. It is the time and uncertainty between a question and a reliable answer.", 0.6, 1.45, 7.6, 0.45, 15, MUTED)
items = [("Team memory", "Who owns this API?", CORAL), ("Swagger files", "Which URL is current?", GOLD), ("Testing tools", "How do I authenticate?", CYAN), ("Chat & documents", "What changed since last release?", MINT)]
for i, (a, b, c) in enumerate(items):
    x = 0.75 + (i % 2) * 3.4
    y = 2.25 + (i // 2) * 1.45
    card(s, x, y, 3.0, 1.05, a, b, c, WHITE, 15, 11)
text(s, "BEFORE", 8.35, 1.55, 1.0, 0.25, 10, CORAL, True)
for i, label in enumerate(["Teams", "Swagger", "Postman", "Excel", "Chat", "Memory"]):
    flow_node(s, label, 8.35 + (i % 2) * 2.05, 2.05 + (i // 2) * 0.95, 1.65, 0.48, PALE_CORAL, CORAL, 11)
    if i > 0:
        line(s, 9.17 + ((i-1) % 2) * 2.05, 2.53 + ((i-1) // 2) * 0.95, 9.17 + (i % 2) * 2.05, 2.53 + (i // 2) * 0.95, CORAL, 1.2, True)
text(s, "The same API question is answered repeatedly, by different people, with different levels of confidence.", 8.35, 5.4, 4.1, 0.7, 15, INK, True, DISPLAY)
add_notes(s, "Say: The product exists because API information is distributed across people, Swagger locations, testing tools, and informal documents.\nKey message: fragmentation creates operational friction and uncertainty.\nProblem addressed: slow discovery, inconsistent testing, and unclear change history.\nBenefit to emphasize: the platform removes repeated searching and manual handoffs.")

# 3
s = base("One operating view for the API landscape", "The solution", 3)
text(s, "TWK Endpoints Hub connects ownership, structure, definitions, testing, change, and standards.", 0.6, 1.45, 7.7, 0.45, 15, MUTED)
flow = [("Teams", CYAN), ("Systems", MINT), ("Environments", GOLD), ("Swagger", CORAL), ("Endpoints", CYAN), ("Testing", MINT), ("Governance", CORAL)]
for i, (label, color) in enumerate(flow):
    x = 0.75 + i * 1.72
    flow_node(s, label, x, 2.55, 1.42, 0.74, WHITE, color, 10)
    if i < len(flow) - 1:
        line(s, x + 1.42, 2.92, x + 1.68, 2.92, MUTED, 1.3, True)
card(s, 0.8, 4.45, 3.65, 1.18, "Find the right API", "Search by system, environment, tag, or endpoint details.", CYAN, PALE_CYAN)
card(s, 4.85, 4.45, 3.65, 1.18, "Act with confidence", "Use reusable auth profiles and run requests from the catalog.", MINT, PALE_MINT)
card(s, 8.9, 4.45, 3.65, 1.18, "See and improve", "Track changes, share knowledge, and validate standards.", CORAL, PALE_CORAL)
add_notes(s, "Say: The solution is a connected operating view, not just another documentation page.\nKey message: every important API action starts from shared context.\nProblem addressed: disconnected tools and stale knowledge.\nBenefit to emphasize: faster, more consistent decisions across teams.")

# 4
s = base("The product model makes ownership visible", "How it works", 4)
text(s, "A simple hierarchy answers the questions people ask first: who owns it, where does it run, and what does it expose?", 0.6, 1.45, 9.3, 0.45, 15, MUTED)
levels = [("TEAM", "Payments", CYAN), ("SYSTEM", "Payment API", MINT), ("ENVIRONMENT", "Production", GOLD), ("SOURCE", "Swagger v2", CORAL), ("ENDPOINT", "GET /api/payments/{id}", CYAN)]
for i, (kind, value, color) in enumerate(levels):
    x = 0.75 + i * 2.45
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.6, 2.1, 1.25, WHITE, WHITE)
    text(s, kind, x + 0.15, 2.82, 1.8, 0.2, 9, color, True)
    text(s, value, x + 0.15, 3.14, 1.8, 0.42, 16, INK, True, DISPLAY)
    if i < len(levels) - 1:
        line(s, x + 2.1, 3.22, x + 2.36, 3.22, color, 2, True)
text(s, "Ownership", 1.0, 4.65, 1.4, 0.25, 11, CYAN, True)
text(s, "Environment", 3.45, 4.65, 1.4, 0.25, 11, GOLD, True)
text(s, "Definition", 5.9, 4.65, 1.4, 0.25, 11, CORAL, True)
text(s, "Action", 8.35, 4.65, 1.4, 0.25, 11, MINT, True)
text(s, "The hierarchy turns an endpoint from an isolated URL into an owned, contextual asset.", 1.0, 5.25, 10.6, 0.48, 20, NAVY, True, DISPLAY, PP_ALIGN.CENTER)
add_notes(s, "Say: This hierarchy is the product's organizing principle.\nKey message: context travels with the endpoint.\nProblem addressed: unclear ownership and environment confusion.\nBenefit to emphasize: teams can locate the right API and understand its operating context quickly.")

# 5
s = base("Discovery moves from asking around to searching once", "Capability", 5)
text(s, "Endpoint discovery combines catalog structure with rich metadata and filters.", 0.6, 1.45, 8.1, 0.45, 15, MUTED)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 2.1, 7.25, 3.65, WHITE, WHITE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 2.42, 6.65, 0.55, PALE_CYAN, PALE_CYAN)
text(s, "Search endpoints, systems, tags...", 1.25, 2.58, 4.9, 0.2, 11, MUTED)
pill(s, "Production", 1.0, 3.25, 1.1, PALE_GOLD, INK)
pill(s, "Payments", 2.25, 3.25, 1.05, PALE_MINT, INK)
pill(s, "GET", 3.45, 3.25, 0.7, PALE_CYAN, INK)
for i, (method, path, status) in enumerate([("GET", "/api/payments/{id}", "Production"), ("POST", "/api/payments/refund", "Test"), ("GET", "/api/payments", "Staging")]):
    y = 3.85 + i * 0.55
    text(s, method, 1.05, y, 0.55, 0.22, 10, CYAN if method == "GET" else CORAL, True)
    text(s, path, 1.75, y, 3.25, 0.22, 11, INK, True, FONT)
    text(s, status, 5.35, y, 1.2, 0.22, 10, MUTED)
text(s, "Search", 8.85, 2.25, 1.1, 0.25, 11, CYAN, True)
text(s, "→", 10.05, 2.25, 0.4, 0.25, 18, MUTED, True, DISPLAY, PP_ALIGN.CENTER)
text(s, "Filter", 10.55, 2.25, 1.1, 0.25, 11, MINT, True)
text(s, "→", 11.75, 2.25, 0.4, 0.25, 18, MUTED, True, DISPLAY, PP_ALIGN.CENTER)
text(s, "Act", 12.25, 2.25, 0.65, 0.25, 11, CORAL, True)
card(s, 8.75, 3.2, 3.55, 1.15, "Rich endpoint detail", "Parameters, headers, request bodies, responses, examples, and deprecation state.", CORAL, PALE_CORAL, 14, 10)
card(s, 8.75, 4.65, 3.55, 1.15, "A shared answer", "The endpoint catalog reduces dependency on personal memory and chat history.", CYAN, PALE_CYAN, 14, 10)
add_notes(s, "Say: Discovery is a direct user value: search, filter, open, understand.\nKey message: users can find an endpoint without knowing which person or repository to ask.\nProblem addressed: API location and metadata uncertainty.\nBenefit to emphasize: less time spent searching and fewer wrong-environment mistakes.")

# 6
s = base("Swagger synchronization keeps the catalog alive", "Capability", 6)
text(s, "The sync engine turns an API definition into a maintained inventory and a visible change signal.", 0.6, 1.45, 9.2, 0.45, 15, MUTED)
steps = [("Swagger URL", CYAN), ("Parse", MINT), ("Compare", GOLD), ("Update catalog", CORAL), ("Notify", CYAN)]
for i, (label, color) in enumerate(steps):
    x = 0.75 + i * 2.4
    flow_node(s, label, x, 2.35, 1.75, 0.7, WHITE, color, 12)
    if i < 4:
        line(s, x + 1.75, 2.7, x + 2.28, 2.7, MUTED, 1.4, True)
for i, (label, color, desc) in enumerate([("NEW", CYAN, "new endpoint"), ("MODIFIED", GOLD, "definition changed"), ("DELETED", CORAL, "no longer present"), ("UNCHANGED", MINT, "no action needed")]):
    x = 0.85 + i * 3.05
    pill(s, label, x, 4.0, 1.25, color, NAVY)
    text(s, desc, x, 4.5, 2.0, 0.25, 11, MUTED)
card(s, 0.85, 5.25, 5.55, 0.82, "Scheduled operation", "Hangfire-backed recurring sync with status and error logs.", CYAN, PALE_CYAN, 14, 10)
card(s, 6.8, 5.25, 5.55, 0.82, "Operational value", "Less manual documentation work and earlier visibility into API drift.", CORAL, PALE_CORAL, 14, 10)
add_notes(s, "Say: Swagger sync is one of the core automation capabilities.\nKey message: the catalog is refreshed from source definitions and changes are classified.\nProblem addressed: outdated documentation and manual comparisons.\nBenefit to emphasize: more accurate inventory and earlier release awareness.")

# 7
s = base("Testing is part of the catalog, not a separate ritual", "Capability", 7)
text(s, "A user can move from endpoint context to an executed request without rebuilding the setup elsewhere.", 0.6, 1.45, 9.5, 0.45, 15, MUTED)
request = [("Endpoint", "GET /api/payments/{id}", CYAN), ("Environment", "Production", GOLD), ("Auth profile", "Payment bearer token", MINT), ("Parameters", "id = 4821", CORAL), ("Execute", "Response + history", CYAN)]
for i, (label, value, color) in enumerate(request):
    y = 2.15 + i * 0.68
    text(s, label.upper(), 0.9, y, 1.45, 0.2, 9, color, True)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.35, y - 0.06, 3.5, 0.38, WHITE, WHITE)
    text(s, value, 2.55, y + 0.03, 3.1, 0.2, 11, INK, True)
    if i < 4:
        line(s, 4.1, y + 0.42, 4.1, y + 0.58, MUTED, 1.1, True)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.0, 2.15, 5.25, 3.55, NAVY, NAVY)
text(s, "200 OK", 7.4, 2.55, 1.3, 0.35, 22, MINT, True, DISPLAY)
text(s, "response time 184 ms", 10.2, 2.64, 1.55, 0.2, 10, RGBColor(178, 197, 208), False, FONT, PP_ALIGN.RIGHT)
text(s, '{\n  "paymentId": 4821,\n  "status": "captured",\n  "amount": 125.00\n}', 7.4, 3.35, 3.9, 1.35, 15, WHITE, False, "Cascadia Mono")
text(s, "Request history preserves what was run and when.", 7.4, 5.15, 3.9, 0.3, 11, CYAN, True)
add_notes(s, "Say: The request runner closes the gap between understanding an API and validating it.\nKey message: endpoint context, environment, authentication, execution, and response live in one flow.\nProblem addressed: repetitive setup in separate tools.\nBenefit to emphasize: faster verification with a reusable request history.")

# 8
s = base("Reusable authentication removes repetitive setup", "Capability", 8)
text(s, "Auth Profiles support the real-world variety of internal APIs while keeping request execution consistent.", 0.6, 1.45, 9.5, 0.45, 15, MUTED)
auths = [("No Auth", "public / internal", CYAN), ("Bearer Token", "static token", MINT), ("API Key", "header-based", GOLD), ("Basic Auth", "user + password", CORAL), ("Custom Header", "arbitrary header", CYAN), ("Auto Login", "login → extract → inject", MINT)]
for i, (name, desc, color) in enumerate(auths):
    x = 0.8 + (i % 3) * 4.05
    y = 2.2 + (i // 3) * 1.25
    card(s, x, y, 3.55, 0.9, name, desc, color, WHITE, 14, 10)
text(s, "Auth profiles are reusable, testable, and stored with encrypted secrets in the implementation.", 0.9, 5.25, 11.0, 0.42, 18, NAVY, True, DISPLAY, PP_ALIGN.CENTER)
add_notes(s, "Say: APIs do not share one authentication pattern, so the product supports six implemented profile types.\nKey message: authentication becomes reusable configuration instead of repeated manual work.\nProblem addressed: token and header setup friction.\nBenefit to emphasize: quicker testing with consistent handling of credentials.\nDo not claim broader security certification; the code evidence supports encrypted secret storage, not a compliance certification.")

# 9
s = base("Change history turns drift into a conversation", "Capability", 9)
text(s, "Instead of comparing files manually, users can see what changed on an endpoint and when.", 0.6, 1.45, 8.8, 0.45, 15, MUTED)
text(s, "BEFORE", 0.95, 2.25, 1.0, 0.2, 10, MUTED, True)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 2.6, 4.8, 1.45, PALE_CYAN, PALE_CYAN)
text(s, "GET /orders/{id}", 1.25, 3.08, 3.9, 0.35, 22, INK, True, "Cascadia Mono")
text(s, "schema snapshot · 12 Jun", 1.25, 3.58, 3.5, 0.2, 10, MUTED)
text(s, "MODIFIED", 6.15, 2.25, 1.3, 0.2, 10, GOLD, True)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.1, 2.6, 5.35, 1.45, PALE_GOLD, PALE_GOLD)
text(s, "GET /orders/{orderId}", 6.45, 3.08, 4.5, 0.35, 22, INK, True, "Cascadia Mono")
text(s, "schema snapshot · 18 Jun", 6.45, 3.58, 3.5, 0.2, 10, MUTED)
line(s, 5.75, 3.3, 6.0, 3.3, CORAL, 2, True)
card(s, 1.0, 4.9, 3.5, 0.9, "What changed", "Path, method, metadata, and schema snapshots are available to inspect.", CORAL, PALE_CORAL, 14, 10)
card(s, 4.9, 4.9, 3.5, 0.9, "Who needs to know", "Teams can react to change without rediscovering the API context.", CYAN, PALE_CYAN, 14, 10)
card(s, 8.8, 4.9, 3.5, 0.9, "Why it matters", "Release awareness becomes visible and traceable.", MINT, PALE_MINT, 14, 10)
add_notes(s, "Say: Change history is a management-friendly capability because it makes API evolution visible.\nKey message: changes are recorded as an understandable before-and-after story.\nProblem addressed: silent or manually discovered API drift.\nBenefit to emphasize: better release communication and fewer surprises for consumers.")

# 10
s = base("Collaboration is built around the endpoint", "Capability", 10)
text(s, "Favorites, sharing, and real-time notifications create a lightweight handoff between API owners and consumers.", 0.6, 1.45, 10.0, 0.45, 15, MUTED)
flow_node(s, "Team A\nowns API", 0.9, 2.5, 2.1, 1.0, PALE_CYAN, CYAN, 14)
line(s, 3.05, 3.0, 4.0, 3.0, CYAN, 2, True)
flow_node(s, "Share endpoint\nwith Team B", 4.05, 2.5, 2.25, 1.0, PALE_MINT, MINT, 14)
line(s, 6.35, 3.0, 7.3, 3.0, MINT, 2, True)
flow_node(s, "Notification\narrives", 7.35, 2.5, 2.0, 1.0, PALE_GOLD, GOLD, 14)
line(s, 9.45, 3.0, 10.3, 3.0, GOLD, 2, True)
flow_node(s, "Team B\ntests it", 10.35, 2.5, 1.9, 1.0, PALE_CORAL, CORAL, 14)
for i, (label, body, color) in enumerate([("Favorites", "keep important endpoints close", CYAN), ("Sharing", "target users or teams", MINT), ("Notifications", "real-time signal + unread state", CORAL)]):
    card(s, 1.0 + i * 4.05, 4.5, 3.55, 0.95, label, body, color, WHITE, 14, 10)
add_notes(s, "Say: Collaboration does not need a separate project or message thread. It is attached to the endpoint.\nKey message: sharing creates a direct path from ownership to consumption.\nProblem addressed: API handoffs lost in chat or personal lists.\nBenefit to emphasize: faster adoption and better visibility of shared assets.")

# 11
s = base("A lightweight escape hatch for fast validation", "Capability", 11)
text(s, "Embedded Tests support arbitrary HTTP requests when a user needs to validate a URL before building a full catalog path.", 0.6, 1.45, 10.8, 0.45, 15, MUTED)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 2.15, 5.9, 3.55, WHITE, WHITE)
for i, (label, value) in enumerate([("METHOD", "POST"), ("URL", "https://api.example.test/health"), ("TOKEN", "••••••••••••"), ("BODY", '{ "check": "deep" }')]):
    y = 2.5 + i * 0.68
    text(s, label, 1.1, y, 0.9, 0.2, 9, CORAL if label == "URL" else CYAN, True)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.15, y - 0.05, 4.1, 0.38, PAPER, PAPER)
    text(s, value, 2.35, y + 0.03, 3.7, 0.2, 11, INK, label == "METHOD", "Cascadia Mono" if label == "BODY" else FONT)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.55, 2.15, 4.7, 3.55, NAVY, NAVY)
text(s, "200 OK", 7.95, 2.65, 1.5, 0.35, 22, MINT, True, DISPLAY)
text(s, "Signed request supported for HMAC scenarios", 7.95, 3.25, 3.6, 0.45, 14, WHITE, True, DISPLAY)
text(s, "Save and reuse test cases when a quick check becomes a repeatable workflow.", 7.95, 4.35, 3.5, 0.5, 12, RGBColor(196, 213, 222))
add_notes(s, "Say: Embedded Tests are useful for fast validation and edge cases that are not yet represented in the endpoint catalog.\nKey message: the platform supports both structured catalog testing and quick exploratory testing.\nProblem addressed: waiting for full setup before a basic API check.\nBenefit to emphasize: shorter feedback loops during troubleshooting and integration work.")

# 12
s = base("Governance turns standards into executable rules", "Capability", 12)
text(s, "The governance pipeline is explicit: define standards, assign scope, validate endpoints, and report results.", 0.6, 1.45, 10.4, 0.45, 15, MUTED)
steps = [("Define\nstandard", CYAN), ("Create\nrules", MINT), ("Assign\ntarget", GOLD), ("Run\nvalidation", CORAL), ("Report\ncompliance", CYAN)]
for i, (label, color) in enumerate(steps):
    x = 0.75 + i * 2.42
    flow_node(s, label, x, 2.35, 1.78, 0.9, WHITE, color, 13)
    if i < 4:
        line(s, x + 1.78, 2.8, x + 2.29, 2.8, MUTED, 1.4, True)
text(s, "Implemented rule types", 0.85, 4.05, 2.3, 0.25, 11, NAVY, True)
for i, label in enumerate(["Regex", "Required field", "Response schema", "HTTP method", "Status code", "Custom expression", "Pagination"]):
    x = 0.9 + (i % 4) * 2.85
    y = 4.55 + (i // 4) * 0.5
    pill(s, label, x, y, 2.35, PALE_CYAN if i % 2 == 0 else PALE_MINT, INK)
add_notes(s, "Say: Governance is not a static policy document; the implementation models standards, rules, assignments, validation, and results.\nKey message: organizational expectations can become repeatable checks.\nProblem addressed: inconsistent API quality and manual compliance review.\nBenefit to emphasize: measurable standards that can be applied at the right organizational scope.")

# 13
s = base("Compliance reports make quality actionable", "Capability", 13)
text(s, "Validation results provide a score, severity context, affected endpoints, and suggested fixes.", 0.6, 1.45, 9.0, 0.45, 15, MUTED)
shape(s, MSO_SHAPE.OVAL, 0.95, 2.25, 2.1, 2.1, PALE_MINT, PALE_MINT)
text(s, "87%", 1.28, 2.85, 1.45, 0.42, 30, MINT, True, DISPLAY, PP_ALIGN.CENTER)
text(s, "compliance score", 1.23, 3.42, 1.55, 0.2, 10, INK, True, FONT, PP_ALIGN.CENTER)
for i, (label, value, color) in enumerate([("Info", "5", CYAN), ("Warning", "12", GOLD), ("Error", "3", CORAL), ("Critical", "1", NAVY)]):
    y = 2.25 + i * 0.68
    pill(s, label, 4.05, y, 1.15, color, WHITE if color in [CORAL, NAVY] else INK)
    text(s, value, 5.5, y + 0.06, 0.45, 0.2, 15, INK, True, DISPLAY)
    text(s, "violations", 6.05, y + 0.08, 1.1, 0.2, 10, MUTED)
card(s, 8.05, 2.25, 4.25, 1.0, "Endpoint summary", "Which endpoints failed, and which rule caused it?", CORAL, PALE_CORAL, 14, 10)
card(s, 8.05, 3.55, 4.25, 1.0, "Suggested fix", "What should the owner change next?", MINT, PALE_MINT, 14, 10)
card(s, 8.05, 4.85, 4.25, 1.0, "Scope and history", "Run against the right team, system, environment, or source.", CYAN, PALE_CYAN, 14, 10)
add_notes(s, "Say: Compliance becomes useful when it points to a specific endpoint, rule, severity, and next action.\nKey message: reporting supports remediation, not just scoring.\nProblem addressed: standards that are difficult to measure or act on.\nBenefit to emphasize: clearer prioritization for API owners and governance teams.\nThe 87% figure is illustrative for the visual; the implementation supports score reporting but does not establish an organizational benchmark.")

# 14
s = base("The dashboard gives management a live operating view", "Visibility", 14)
text(s, "The dashboard brings inventory, change signals, broken sources, and recent sync activity into one surface.", 0.6, 1.45, 10.2, 0.45, 15, MUTED)
metrics = [("Teams", "12", CYAN), ("Systems", "34", MINT), ("Environments", "86", GOLD), ("Endpoints", "1,248", CORAL)]
for i, (label, value, color) in enumerate(metrics):
    x = 0.8 + i * 3.05
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.15, 2.55, 1.0, WHITE, WHITE)
    text(s, value, x + 0.2, 2.38, 1.4, 0.36, 24, color, True, DISPLAY)
    text(s, label, x + 0.2, 2.82, 1.5, 0.2, 10, MUTED)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 3.55, 7.0, 1.75, WHITE, WHITE)
text(s, "Recent sync activity", 1.1, 3.8, 2.0, 0.25, 14, INK, True, DISPLAY)
for i, (name, status, color) in enumerate([("Payment API / Production", "Success", MINT), ("Orders API / Test", "Modified", GOLD), ("Identity API / Staging", "Broken URL", CORAL)]):
    y = 4.25 + i * 0.32
    text(s, name, 1.1, y, 3.3, 0.18, 10, INK)
    text(s, status, 5.65, y, 1.2, 0.18, 10, color, True)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.15, 3.55, 4.35, 1.75, NAVY, NAVY)
text(s, "CHANGE SIGNAL", 8.5, 3.85, 1.7, 0.2, 9, CYAN, True)
text(s, "+18 new  ·  6 modified  ·  2 deleted", 8.5, 4.25, 3.5, 0.3, 17, WHITE, True, DISPLAY)
text(s, "A management view of API inventory health and movement.", 8.5, 4.75, 3.35, 0.25, 10, RGBColor(191, 208, 217))
add_notes(s, "Say: The dashboard is the management lens over the operational catalog.\nKey message: inventory and movement are visible without opening individual systems.\nProblem addressed: lack of centralized API landscape visibility.\nBenefit to emphasize: quicker prioritization and earlier awareness of sync problems or change.")

# 15
s = base("The platform supports different stakeholders with one shared model", "Organizational value", 15)
roles = [("Developers", "discover, run, generate", CYAN), ("API owners", "publish, share, track", MINT), ("Architects", "see structure, standards", GOLD), ("QA", "validate behavior", CORAL), ("Governance", "rules, scores, findings", CYAN), ("Management", "inventory, movement, risk", MINT)]
for i, (role, value, color) in enumerate(roles):
    x = 0.8 + (i % 3) * 4.1
    y = 2.0 + (i // 3) * 1.75
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.55, 1.22, WHITE, WHITE)
    shape(s, MSO_SHAPE.OVAL, x + 0.22, y + 0.27, 0.68, 0.68, color, color)
    text(s, str(i + 1), x + 0.22, y + 0.43, 0.68, 0.2, 14, NAVY, True, DISPLAY, PP_ALIGN.CENTER)
    text(s, role, x + 1.1, y + 0.25, 2.15, 0.27, 15, INK, True, DISPLAY)
    text(s, value, x + 1.1, y + 0.64, 2.15, 0.2, 10, MUTED)
text(s, "One catalog, multiple decisions: build, consume, assure, govern, and prioritize.", 0.9, 5.85, 11.5, 0.35, 19, NAVY, True, DISPLAY, PP_ALIGN.CENTER)
add_notes(s, "Say: Different roles use different slices of the same product, which is what makes centralization valuable.\nKey message: the product connects delivery, quality, ownership, and management visibility.\nProblem addressed: role-specific information silos.\nBenefit to emphasize: a shared operating language across the API lifecycle.")

# 16
s = base("A new API can move from release to governed adoption", "End-to-end scenario", 16, True)
text(s, "Example: a new Payment API is released and becomes usable across the organization.", 0.65, 1.45, 9.0, 0.4, 15, RGBColor(193, 211, 221))
scenario = ["Team owns\nPayment API", "Swagger\nsyncs", "Endpoint\nappears", "Developer\nfinds it", "Auth +\ntest", "Share with\nconsumer", "Change is\ntracked", "Governance\nvalidates"]
for i, label in enumerate(scenario):
    x = 0.6 + (i % 4) * 3.1
    y = 2.25 + (i // 4) * 1.55
    color = [CYAN, MINT, GOLD, CORAL][i % 4]
    flow_node(s, label, x, y, 2.2, 0.82, NAVY, color, 13)
    if i in [0, 1, 2, 4, 5, 6]:
        line(s, x + 2.2, y + 0.41, x + 2.92, y + 0.41, color, 1.4, True)
text(s, "The story is continuous: source → visibility → action → collaboration → assurance.", 1.0, 5.85, 11.2, 0.35, 20, WHITE, True, DISPLAY, PP_ALIGN.CENTER)
add_notes(s, "Say: This is the end-to-end story management should remember. A release becomes discoverable, testable, shareable, traceable, and governable.\nKey message: the value comes from the connected flow, not one isolated screen.\nProblem addressed: handoffs that lose context between release and adoption.\nBenefit to emphasize: a shorter path from API release to confident organizational use.")

# 17
s = base("What is implemented today, and what remains open", "Product maturity", 17)
text(s, "The current codebase supports a broad working platform. The remaining gaps are opportunities, not current claims.", 0.6, 1.45, 10.3, 0.45, 15, MUTED)
card(s, 0.8, 2.15, 5.75, 3.55, "Implemented capabilities", "Authentication and RBAC\nTeams → systems → environments\nEndpoint catalog and rich metadata\nSwagger sync, schedules, and change logs\nRequest runner + six auth profile types\nFavorites, sharing, notifications\nCode generation and export\nGovernance rules, validation, compliance reports\nDashboard, embedded tests, bilingual UI", MINT, PALE_MINT, 17, 12)
card(s, 6.8, 2.15, 5.75, 3.55, "Future opportunities", "Automated unit / E2E test coverage\nUsage analytics and API performance views\nRetention, backup, and bulk data management\nCI/CD, gateway, Slack, Teams, or Jira integrations\nMore background job types\nWebhooks and advanced lifecycle management\nBreaking-change and deprecation workflows", CORAL, PALE_CORAL, 17, 12)
text(s, "Current capability claims are grounded in frontend and backend evidence; future items are clearly separated.", 1.0, 6.2, 11.2, 0.3, 13, NAVY, True, DISPLAY, PP_ALIGN.CENTER)
add_notes(s, "Say: Credibility matters, so this presentation separates working capabilities from future potential.\nKey message: the product is already a substantial internal platform, while the roadmap remains open.\nProblem addressed: overclaiming maturity or presenting ideas as delivered features.\nBenefit to emphasize: stakeholders can evaluate the current value and discuss the next investment areas clearly.")

# 18
s = base("From fragmented API knowledge to organizational leverage", "Executive close", 18, True)
text(s, "TWK Endpoints Hub makes APIs easier to find, easier to use, easier to change, and easier to govern.", 0.7, 1.65, 8.9, 0.65, 28, WHITE, True, DISPLAY)
finals = [("01", "Discover APIs from one place", CYAN), ("02", "Keep the catalog synchronized", MINT), ("03", "Test with reusable context", GOLD), ("04", "Track changes and collaborate", CORAL), ("05", "Apply measurable standards", CYAN)]
for i, (n, label, color) in enumerate(finals):
    y = 3.0 + i * 0.55
    text(s, n, 0.85, y, 0.42, 0.2, 10, color, True)
    text(s, label, 1.5, y - 0.02, 5.4, 0.25, 15, WHITE, True, DISPLAY)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.2, 2.75, 3.85, 2.5, NAVY, CYAN)
text(s, "THE VALUE", 8.6, 3.15, 1.6, 0.2, 10, CYAN, True)
text(s, "Less searching.\nLess repetition.\nMore visibility.\nMore confidence.", 8.6, 3.65, 2.8, 1.2, 22, WHITE, True, DISPLAY)
text(s, "TWK Endpoints Hub", 0.7, 6.55, 3.5, 0.25, 12, RGBColor(180, 200, 210), True)
add_notes(s, "Say: Close by returning to the business story. TWK Endpoints Hub centralizes API knowledge and turns it into faster action and clearer governance.\nKey message: this is organizational leverage built on practical developer workflows.\nProblem addressed: fragmented knowledge, manual effort, and limited visibility.\nBenefit to emphasize: a common platform for discovery, delivery, collaboration, and control.")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Created {OUT}")
